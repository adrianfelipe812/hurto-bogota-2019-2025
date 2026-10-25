from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(__file__).parent / 'Decisiones_de_diseno_visual_hurto_Bogota.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.72); sec.bottom_margin = Inches(.68)
sec.left_margin = Inches(.82); sec.right_margin = Inches(.82)

NAVY='173249'; TEAL='168B80'; PALE='EAF1F4'; GRAY='536675'; BORDER='D9E1E5'
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10); styles['Normal'].font.color.rgb=RGBColor.from_string('263746')
styles['Normal'].paragraph_format.space_after=Pt(5)
for sname,size in [('Title',25),('Heading 1',17),('Heading 2',12),('Heading 3',10)]:
    s=styles[sname]; s.font.name='Aptos Display'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_before=Pt(11); s.paragraph_format.space_after=Pt(5)

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def borders(table):
    pr=table._tbl.tblPr; b=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'5'); e.set(qn('w:color'),BORDER); b.append(e)
    pr.append(b)
def cellpad(cell, top=100, start=110, bottom=100, end=110):
    tc=cell._tc; pr=tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
    for edge,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        el=OxmlElement('w:'+edge); el.set(qn('w:w'),str(val)); el.set(qn('w:type'),'dxa'); mar.append(el)
    pr.append(mar)
def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths): c.width=Inches(w)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,NAVY); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; cellpad(c)
        for r in c.paragraphs[0].runs: r.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
    trPr=t.rows[0]._tr.get_or_add_trPr(); rep=OxmlElement('w:tblHeader'); trPr.append(rep)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        if widths:
            for c,w in zip(t.columns,widths): c.width=Inches(w)
        for i,val in enumerate(row):
            cells[i].text=val; cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; cellpad(cells[i])
            if ri%2==1: shade(cells[i],'F4F7F8')
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(1)
                for r in p.runs: r.font.size=Pt(9)
    borders(t); doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t
def bullet(text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2); p.add_run(text); return p
def labeled(label, text):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4)
    r=p.add_run(label+' '); r.bold=True; p.add_run(text); return p

# Encabezado y pie discretos
header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=header.add_run('PROYECTO COLABORATIVO · VISUALIZACIÓN DE DATOS'); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('Hurto a personas en Bogotá · Decisiones de diseño'); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)

p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Decisiones de diseño visual\nHurto a personas en Bogotá')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Documento colaborativo de propósito, método y criterios de calidad'); r.italic=True; r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string(GRAY)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Periodo analizado: enero–diciembre de 2019 a 2025'); r.bold=True; r.font.size=Pt(10); r.font.color.rgb=RGBColor.from_string(TEAL)

doc.add_heading('Síntesis del proyecto', level=1)
doc.add_paragraph('Como equipo, construimos tres visualizaciones complementarias de los casos de hurto a personas registrados en Bogotá: una comparación de distribuciones mensuales por año, un mapa de calor año–mes y una comparación de la evolución mensual de 2023, 2024 y 2025. El análisis usa los doce meses de siete años completos, de modo que las comparaciones temporales se realizan sobre periodos equivalentes. La pieza combina exploración estadística con comunicación explicativa; el tercer gráfico presenta una comparación focalizada de los tres años más recientes disponibles en el conjunto analizado.')

doc.add_heading('1. Formación del equipo y distribución de roles', level=1)
doc.add_paragraph('El trabajo se organiza entre tres integrantes. Los roles se pueden rotar, pero cada versión final debe identificar quién asumió cada responsabilidad y quién hizo la revisión cruzada.')
table(['Rol','Responsabilidad principal','Integrante'],[
('Analista','Revisa estructura, alcance y consistencia de los datos; define agregaciones y comparaciones; interpreta resultados sin atribuir causalidad no demostrada.','[Nombre del integrante 1]'),
('Diseñador/a visual','Selecciona gráficos y codificaciones; implementa la jerarquía visual, etiquetas, paleta y anotaciones.','[Nombre del integrante 2]'),
('Revisor/a de calidad','Contrasta figuras con los datos y el código; revisa escalas, filtros, legibilidad, accesibilidad y reproducibilidad.','[Nombre del integrante 3]'),
], [1.25,4.55,1.55])
labeled('Acuerdo de colaboración:', 'cada integrante revisa al menos una decisión tomada por otra persona. Si los roles rotan, registrar la rotación y la versión revisada en el cuaderno o en la bitácora del equipo.')

doc.add_heading('2. Selección y preparación de los datos', level=1)
doc.add_paragraph('Se utiliza el archivo público Hurto_a_personas.csv, identificado como fuente SIEDCO (Sistema de Información Estadístico, Delincuencial, Contravencional y Operativo). El archivo contiene registros con año, mes, departamento y cantidad. El código valida los campos requeridos, convierte la cantidad a valor numérico entero no negativo y comprueba que los registros incluidos correspondan a Bogotá.')
labeled('Unidad de análisis:', 'total de casos registrados por mes y año, calculado sumando la variable Cantidad en las filas detalladas.')
labeled('Ventana temporal:', 'enero–diciembre de 2019–2025. Se comparan siete años completos y doce meses por año; 2026 no forma parte del análisis actual.')
labeled('Controles implementados:', 'presencia de los siete años, doce meses por año, ausencia de combinaciones año–mes duplicadas después de agregar y conservación del total de Cantidad durante la agregación. No se imputan meses faltantes.')
labeled('Referencia que debe completar el equipo:', 'añadir el enlace público exacto de descarga, la fecha en que el equipo obtuvo el archivo y, si el portal publica metadatos, el nombre oficial del conjunto y su cobertura. La fecha de actualización reportada en las figuras es 21/10/2026; no se debe confundir con la fecha de descarga.')

doc.add_heading('3. Propósito de las visualizaciones', level=1)
doc.add_paragraph('El conjunto tiene un propósito mixto y explícito. Los dos primeros gráficos son principalmente exploratorios: ayudan a encontrar variabilidad, patrones mensuales y diferencias entre años. El tercero es explicativo: comunica la comparación de la trayectoria mensual entre 2023, 2024 y 2025, destacando 2025. “Explicativo” describe la intención de comunicación; no significa que el gráfico explique causas del delito.')

doc.add_heading('4. Audiencia, contexto y uso', level=1)
labeled('Audiencia prevista:', 'estudiantes y lectores con conocimientos básicos de estadística descriptiva y lectura de gráficos; también puede servir a equipos académicos o de análisis de seguridad que necesiten examinar registros agregados.')
labeled('Acción que busca apoyar:', 'facilitar preguntas de seguimiento sobre la distribución mensual, la estacionalidad y los cambios observados entre años. La visualización no basta por sí sola para asignar recursos ni evaluar políticas; esas decisiones requieren contexto adicional, denominadores poblacionales, información operacional y análisis causal apropiado.')
labeled('Conocimiento asumido:', 'comprensión de conteos, meses y años; en el boxplot, familiaridad básica con mediana, cuartiles y rango intercuartílico. Se incluyen etiquetas y anotaciones para que los gráficos sigan siendo legibles sin conocimientos estadísticos avanzados.')

doc.add_heading('5. Selección y justificación de los gráficos', level=1)
doc.add_heading('Gráfico 1 · Distribución mensual por año', level=2)
labeled('Título recomendado:', 'Distribución mensual de casos registrados por año (2019–2025).')
doc.add_paragraph('Se utiliza un boxplot por año acompañado de los doce puntos mensuales. El boxplot resume mediana, cuartiles y dispersión; los puntos preservan las observaciones y permiten apreciar los meses concretos. Las medianas anotadas corresponden a los doce totales mensuales de cada año. Se ocultan los símbolos automáticos de valores atípicos del boxplot para evitar duplicarlos, pero se conservan todos los puntos mensuales; por tanto, ningún mes se excluye del conjunto visible.')

doc.add_heading('Gráfico 2 · Comparación entre meses y años', level=2)
labeled('Título recomendado:', 'Casos registrados por mes y año (2019–2025).')
doc.add_paragraph('El mapa de calor organiza los meses en columnas y los años en filas. La intensidad facilita localizar combinaciones año–mes relativamente altas o bajas, mientras que la cifra anotada conserva la lectura exacta. Una sola escala común se aplica a todas las celdas, sin doble eje. La paleta azul–roja representa magnitud del conteo; el rojo no significa mayor gravedad individual.')

doc.add_heading('Gráfico 3 · Evolución mensual de los tres años recientes', level=2)
labeled('Título recomendado:', 'En 2025, los registros mensuales se mantuvieron por debajo de 2023.')
doc.add_paragraph('Se utiliza un gráfico de líneas para comparar el perfil de enero a diciembre de los tres años más recientes incluidos. La frase del título se limita a la comparación descriptiva visible entre 2025 y 2023; no afirma una disminución continua dentro de 2025 ni atribuye los cambios a una política o evento. 2025 se destaca y 2023–2024 se muestran como referencias. Los puntos marcan los doce valores observados; las curvas PCHIP solo interpolan visualmente entre ellos y no son observaciones adicionales. La interpretación debe anclarse en los marcadores y los conteos originales.')

doc.add_heading('6. Codificación visual', level=1)
table(['Gráfico','Posición y marcas','Color y otros canales'],[
('Boxplot','Año en x; casos mensuales en y. Caja = Q1–Q3; línea = mediana; bigotes según 1,5 × RIC; puntos = los doce meses.','Cajas azul claro; 2025 destacado en teal; puntos oscuros. Color no sustituye la posición ni oculta observaciones.'),
('Mapa de calor','Mes en x; año en y; valor exacto impreso en cada celda.','Intensidad de una paleta divergente azul–roja, con un rango común. La cifra anotada evita depender solo del color.'),
('Líneas','Mes en x; casos registrados en y; puntos en cada mes.','Teal y trazo más grueso para 2025; tonos grises para 2023–2024. Etiquetas directas al final evitan una leyenda separada.'),
],[1.15,3.25,2.95])

doc.add_heading('7. Legibilidad y jerarquía', level=1)
bullet('La pieza consolidada usa el título general “Hurto a personas en Bogotá 2019–2025 (enero–diciembre)”, que declara lugar, tema y periodo.')
bullet('Cada eje identifica la variable y su unidad de conteo; los meses y años siguen orden cronológico y los valores usan separador de miles.')
bullet('Los títulos de panel describen el contenido sin adelantar conclusiones que los datos no sostienen. En una versión de presentación se puede añadir una frase de hallazgo después de verificarla numéricamente.')
bullet('Se reduce carga cognitiva al eliminar marcos superiores y laterales, leyendas redundantes y conectores de anotaciones; se mantienen puntos y etiquetas para no perder información relevante.')
bullet('La fuente y la fecha de actualización aparecen en las figuras. El pie debe acompañarse en la entrega con la referencia o enlace exacto del conjunto público.')

doc.add_heading('8. Color, estilo y accesibilidad', level=1)
doc.add_paragraph('Se mantiene una base tipográfica y visual común. En el boxplot y las líneas se emplean colores categóricos para distinguir años, con énfasis en el año más reciente (2025). En el mapa, la paleta divergente azul–roja hace visibles valores relativamente bajos y altos respecto del rango mostrado; como los datos son conteos no negativos y no existe un punto de referencia sustantivo en cero en el centro, el equipo debe considerar si una paleta secuencial perceptualmente ordenada comunicaría la magnitud con menos ambigüedad. Si se conserva RdBu_r, se documentará como una decisión de contraste, no como una escala con significado positivo/negativo.')
doc.add_paragraph('Para accesibilidad, los valores del mapa también están escritos; las series del gráfico de líneas se distinguen por tono, grosor y marcadores; se conserva contraste suficiente entre texto y fondo. La versión final se revisará en escala de grises o con un simulador de deficiencia de visión cromática si se va a distribuir ampliamente.')

doc.add_heading('9. Chequeo ético e interpretativo', level=1)
bullet('Los ejes verticales de las visualizaciones empiezan en cero. Esto mantiene contexto para los conteos y evita exagerar diferencias mediante recortes.')
bullet('No se usa 3D ni ornamentación que cambie la percepción de las magnitudes.')
bullet('Los gráficos 1 y 2 incluyen todos los años 2019–2025 y los doce meses. El gráfico 3 limita intencionalmente la selección a 2023–2025 para facilitar la comparación reciente; el recorte se declara en el título.')
bullet('La variable representa casos registrados y agregados, no tasas por población ni necesariamente la totalidad de hechos ocurridos. Diferencias en denuncia, registro o clasificación pueden afectar los conteos.')
bullet('Las asociaciones temporales son descriptivas. Estos gráficos no identifican causas ni permiten evaluar por sí solos el efecto de intervenciones.')
bullet('El color rojo del mapa codifica valor alto en la escala, no riesgo individual, severidad, culpabilidad ni probabilidad de victimización.')

doc.add_heading('10. Construcción reproducible en Python', level=1)
doc.add_paragraph('El proyecto se implementa en Python con pandas para lectura y agregación, Seaborn para prototipar los gráficos estadísticos y Matplotlib para ajustar composición, títulos, ejes, anotaciones y exportación. El cuaderno y el script deben ejecutarse desde la carpeta del proyecto, con Hurto_a_personas.csv disponible en ella.')
table(['Componente','Implementación del proyecto'],[
('Prototipado','sns.boxplot y sns.stripplot para distribución; sns.heatmap para matriz año–mes; Matplotlib para las series temporales con valores mensuales.'),
('Acabados','Tema común de Seaborn; tipografía DejaVu Sans; formato de miles; etiquetas directas y anotaciones; eje vertical desde cero.'),
('Reproducibilidad','Script visualizaciones_hurto_ene_sep.py y cuaderno 01_visualizaciones_hurto_ene_sep.ipynb; configuración de semilla fija para el desplazamiento horizontal del stripplot.'),
('Dependencias','pandas, seaborn, matplotlib y scipy, declaradas en requirements.txt.'),
('Exportación actual','Tres PNG separados a 300 dpi en salidas/. La entrega final también debe incluir PDF; actualizar el código para exportar PDF vectorial o generar un PDF que reúna las figuras.'),
],[1.35,5.95])
doc.add_paragraph('Nota de mantenimiento: los nombres del notebook y del script contienen “ene_sep”, pero el código actual analiza enero–diciembre de 2019–2025 y excluye 2026. Para evitar ambigüedad al ejecutar o compartir, conviene renombrar esos archivos o explicar la discrepancia en el README. La figura adjunta previamente mostraba otro corte temporal; esta especificación se ajusta al contenido actual de la carpeta del proyecto.')

doc.add_heading('11. Exportación y paquete de entrega', level=1)
doc.add_paragraph('El paquete colaborativo se considerará completo cuando contenga las figuras finales, este documento de decisiones y el código reproducible, junto con los datos o instrucciones públicas necesarias para obtenerlos.')
table(['Entregable','Estado en la carpeta del proyecto','Criterio de cierre'],[
('Figura 1 · boxplot','PNG a 300 dpi disponible en salidas/.','Confirmar título, resolución y legibilidad al tamaño de uso.'),
('Figura 2 · mapa de calor','PNG a 300 dpi disponible en salidas/.','Revisar rango y elección de paleta; mantener valores legibles.'),
('Figura 3 · líneas','PNG a 300 dpi disponible en salidas/.','Mantener datos mensuales marcados; declarar 2023–2025.'),
('PDF de figuras','Pendiente de exportación.','Incluir las tres figuras en buena resolución o vectoriales, con títulos y fuente.'),
('Documento de decisiones','Este archivo Word.','Completar nombres del equipo y referencia pública exacta del conjunto.'),
('Código reproducible','Script y notebook disponibles; requirements.txt disponible.','Probar la ejecución desde la carpeta y verificar que produce las salidas incluidas.'),
],[1.55,2.45,3.3])

doc.add_heading('Registro final de revisión del equipo', level=1)
table(['Control','Responsable','Resultado / fecha'],[
('Datos, filtros y agregación','[Nombre]','[Aprobado / ajuste pendiente]'),
('Elección de gráficos y codificación','[Nombre]','[Aprobado / ajuste pendiente]'),
('Legibilidad, accesibilidad y ética','[Nombre]','[Aprobado / ajuste pendiente]'),
('Ejecución reproducible y exportación','[Nombre]','[Aprobado / ajuste pendiente]'),
],[2.5,1.7,3.1])

doc.add_heading('Referencias y materiales del proyecto', level=1)
bullet('Fuente de datos: SIEDCO, Sistema de Información Estadístico, Delincuencial, Contravencional y Operativo. Enlace público exacto: [completar por el equipo]. Fecha de descarga: [completar].')
bullet('Archivo de datos: Hurto_a_personas.csv. Fecha de actualización indicada en el proyecto: 21/10/2026.')
bullet('Código: visualizaciones_hurto_ene_sep.py; cuaderno: 01_visualizaciones_hurto_ene_sep.ipynb; dependencias: requirements.txt.')
bullet('Conceptos de diseño aplicados: lectura crítica, elección de marcas y canales visuales, distribuciones, comparación entre grupos, paletas y anotaciones.')

doc.core_properties.title='Decisiones de diseño visual sobre hurto a personas en Bogotá'
doc.core_properties.subject='Proyecto colaborativo de visualización de datos'
doc.core_properties.author='Equipo de trabajo'
doc.save(OUT)
print(OUT)
